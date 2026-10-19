#!/usr/bin/env python3
"""
Builds the individual technical report (Word .docx) for CS7 / S03 from the experiment output.

All numbers in the text, the summary table and the figures come from results/ (written by
run_experiment.py + analyze.py). Static sections (introduction, literature, method, limitations,
future work) are fixed text. Sentences that interpret the data are built from the measured numbers
and branch on what was actually found, so READ THE RESULT PARAGRAPHS BEFORE SUBMITTING and edit
anything that does not match what you saw in the figures.

Usage:  python make_report.py --name "Your Name" --reg 23BCE0434 --batch "B07"
"""
import argparse, re
from pathlib import Path
import numpy as np
import pandas as pd

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except ImportError:
    raise SystemExit("python-docx is needed:  pip install python-docx")

FONT = "Times New Roman"

REFS = [
    'NVIDIA, "CUDA C++ Best Practices Guide," CUDA Toolkit Documentation, v13.0. https://docs.nvidia.com/cuda/archive/13.0.0/cuda-c-best-practices-guide/index.html',
    'S. Mittal and J. S. Vetter, "A Survey of CPU-GPU Heterogeneous Computing Techniques," ACM Computing Surveys, vol. 47, no. 4, Art. 69, 2015. https://doi.org/10.1145/2788396',
    'C.-K. Luk, S. Hong, and H. Kim, "Qilin: Exploiting Parallelism on Heterogeneous Multiprocessors with Adaptive Mapping," in Proc. 42nd IEEE/ACM Int. Symp. on Microarchitecture (MICRO-42), 2009, pp. 45-55.',
    'C. Gregg and K. Hazelwood, "Where is the Data? Why You Cannot Debate CPU vs. GPU Performance Without the Answer," in Proc. IEEE ISPASS, 2011. https://doi.org/10.1109/ISPASS.2011.5762730',
    'R. Nozal and J. L. Bosque, "Exploiting co-execution with oneAPI: heterogeneity from a modern perspective," arXiv:2106.01726, 2021 (Springer version: https://doi.org/10.1007/978-3-030-85665-6_31).',
    'Z. Li, M. Paolieri, and L. Golubchik, "Accelerating Mobile Inference through Fine-Grained CPU-GPU Co-Execution," in Selected Papers of EPEW 2025, LNCS vol. 15657, Springer, 2026, pp. 41-55.',
    'F. Wang et al., "Accurate and Reliable Energy Measurement and Modelling of Data Transfer Between CPU and GPU in Parallel Applications on Heterogeneous Hybrid Platforms," IEEE Trans. Computers, vol. 74, no. 3, pp. 1011-1024, 2025. https://doi.org/10.1109/TC.2024.3504262',
]


# ------------------------------------------------------------------ small formatting helpers
def ms(x):
    if x >= 100:
        return f"{x:.0f} ms"
    if x >= 10:
        return f"{x:.1f} ms"
    return f"{x:.2f} ms"


def pc(x, d=0):
    return f"{100 * x:.{d}f}%"


def sx(x):
    return f"{x:.2f}x" if x < 10 else f"{x:.1f}x"


def join_and(items):
    items = list(items)
    if len(items) <= 1:
        return "".join(items)
    return ", ".join(items[:-1]) + " and " + items[-1]


def nums(vals):
    return join_and(str(int(v)) for v in vals)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_font(run, size=None, bold=None, italic=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_runs(par, text, size=9.5):
    """Supports **bold** segments."""
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if chunk:
            r = par.add_run(chunk)
            set_font(r, size, bold=(i % 2 == 1))


def para(doc, text, size=9.5, after=3, align=None, italic=False, before=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after, pf.space_before, pf.line_spacing = Pt(after), Pt(before), 1.0
    if align == "j":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(p, text, size)
    if italic:
        for r in p.runs:
            r.italic = True
    return p


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after, p.paragraph_format.keep_with_next = Pt(5), Pt(2), True
    r = p.add_run(text)
    set_font(r, 10.5, bold=True)
    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    # real outline level so the Word navigation pane works
    pPr = p._p.get_or_add_pPr()
    lvl = OxmlElement("w:outlineLvl")
    lvl.set(qn("w:val"), "0")
    pPr.append(lvl)
    return p


def caption(doc, text):
    p = para(doc, text, size=8, after=4, italic=True, align="c")
    return p


def table(doc, header, rows, widths=None, size=7.5):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        set_font(r, size, bold=True)
        shade(c, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            set_font(r, size)
    for row in t.rows:
        for i, c in enumerate(row.cells):
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
            if widths:
                c.width = Inches(widths[i])
    for row in t.rows[:-1]:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True
    return t


# ------------------------------------------------------------------ facts from the data
def read_sysinfo(path):
    d = {}
    if Path(path).exists():
        for line in open(path):
            if ":" in line:
                k, v = line.split(":", 1)
                d[k.strip()] = v.strip()
    return d


def verdict(ok, partly=False):
    return "supported" if ok else ("partly supported" if partly else "not supported")


def compute_facts(raw, med, opt):
    f = {}
    f["n_cfg"] = len(opt)
    f["tiles_list"] = sorted(int(v) for v in opt.tiles.unique())
    f["passes_list"] = sorted(int(v) for v in opt.passes.unique())
    f["threads_list"] = sorted(int(v) for v in opt.threads_req.unique())
    f["tile_dim"] = int(raw.tile_dim.iloc[0])
    f["reps"] = int(med.n_reps.max())
    f["t_max"], f["th_max"], f["th_min"] = max(f["tiles_list"]), max(f["threads_list"]), min(f["threads_list"])
    f["p_min"], f["p_max"] = min(f["passes_list"]), max(f["passes_list"])
    opt = opt.copy()
    opt["opt_mid"] = (opt.near_opt_lo + opt.near_opt_hi) / 2
    f["opt"] = opt
    H = opt[(opt.tiles == f["t_max"]) & (opt.threads_req == f["th_max"])].sort_values("passes")
    f["H"] = H
    f["interior"] = int(((opt.best_frac > 0) & (opt.best_frac < 1)).sum())
    f["beats_gpu"] = int(opt.hybrid_beats_gpu_only.sum())
    f["beats_single"] = int(opt.hybrid_beats_best_single.sum())
    f["gpu_faster"] = int((opt.faster_single_device == "GPU-only").sum())
    f["gain_min"], f["gain_max"] = opt.speedup_best_vs_gpu_only.min(), opt.speedup_best_vs_gpu_only.max()

    def mean_delta(sel_cols, hi, lo, key):
        ds = []
        for _, g in opt.groupby(sel_cols):
            a, b = g[g[key] == hi], g[g[key] == lo]
            if len(a) and len(b):
                ds.append(a.opt_mid.iloc[0] - b.opt_mid.iloc[0])
        return float(np.mean(ds)) if ds else float("nan")

    f["d_passes"] = mean_delta(["tiles", "threads_req"], f["p_max"], f["p_min"], "passes") if f["p_max"] != f["p_min"] else float("nan")
    f["d_threads"] = mean_delta(["tiles", "passes"], f["th_max"], f["th_min"], "threads_req") if f["th_max"] != f["th_min"] else float("nan")
    f["d_size"] = mean_delta(["passes", "threads_req"], f["t_max"], min(f["tiles_list"]), "tiles") if f["t_max"] != min(f["tiles_list"]) else float("nan")
    f["pred_in_range"] = int(((opt.predicted_frac >= opt.near_opt_lo - 0.05) & (opt.predicted_frac <= opt.near_opt_hi + 0.05)).sum())
    f["pred_mae"] = float((opt.predicted_frac - opt.best_frac).abs().mean())
    f["pe_median"] = float(opt.partition_eff_best.median())
    f["sync_rel"] = float((opt.sync_ms_at_best / opt.best_ms).median())
    f["sched_rel"] = float((opt.sched_ms_at_best / opt.best_ms).median())
    f["cv_med"], f["cv_max"] = float(med.makespan_cv.median()), float(med.makespan_cv.max())
    f["err_max"] = float(raw.max_rel_err.max())
    f["n_runs"] = len(raw)
    fr = np.sort(med.cpu_frac.unique())
    f["step"] = int(round(100 * float(np.median(np.diff(fr))))) if len(fr) > 1 else 10
    f["chunk"] = nums(sorted(raw.chunk_tiles.unique()))
    f["streams_txt"] = nums(sorted(raw.streams.unique()))
    f["streams_phrase"] = f["streams_txt"] + (" stream" if f["streams_txt"] == "1" else " streams")
    pin = set(raw.pinned.unique())
    f["pinned_txt"] = "pinned" if pin == {1} else ("pageable" if pin == {0} else "pinned and pageable")
    # imbalance at the best split of interior-optimum configs
    imb = []
    for _, r in opt.iterrows():
        if 0 < r.best_frac < 1:
            row = med[(med.tiles == r.tiles) & (med.passes == r.passes) & (med.threads_req == r.threads_req) & (med.cpu_frac == r.best_frac)]
            if len(row):
                imb.append(float(row.imbalance_ms.iloc[0] / row.makespan_ms.iloc[0]))
    f["imb_med"] = float(np.median(imb)) if imb else float("nan")
    f["single_best"] = opt[(opt.best_frac == 0) | (opt.best_frac == 1)]

    def med_row(tiles, passes, th, frac):
        r = med[(med.tiles == tiles) & (med.passes == passes) & (med.threads_req == th) & (med.cpu_frac == frac)]
        return r.iloc[0] if len(r) else None

    util, scal = [], []
    for _, r in H.iterrows():
        m = med_row(f["t_max"], r.passes, f["th_max"], r.best_frac)
        if m is not None:
            util.append(f"passes {int(r.passes)}: CPU busy {100 * m.cpu_util:.0f}%, GPU kernel {100 * m.gpu_kernel_util:.0f}%, GPU active including copies {100 * m.gpu_active:.0f}%")
        lo = opt[(opt.tiles == f["t_max"]) & (opt.passes == r.passes) & (opt.threads_req == f["th_min"])]
        if m is not None and len(lo) and pd.notna(lo.iloc[0].speedup_best_vs_serial) and pd.notna(r.speedup_best_vs_serial) and f["th_max"] != f["th_min"]:
            scal.append(f"passes {int(r.passes)}: {sx(lo.iloc[0].speedup_best_vs_serial)} at {f['th_min']} thread(s) to {sx(r.speedup_best_vs_serial)} at {f['th_max']} (efficiency {m.efficiency:.2f})")
    f["util_txt"], f["scal_txt"] = "; ".join(util), "; ".join(scal)
    nz = med.loc[med.makespan_cv.idxmax()]
    f["noisy"] = (int(nz.passes), int(nz.tiles), int(nz.threads_req), float(nz.cpu_frac), float(nz.makespan_ms), float(nz.makespan_cv))
    return f


def trend_word(d, unit="points"):
    if np.isnan(d):
        return "could not be assessed"
    if d <= -0.05:
        return f"fell by about {abs(100 * d):.0f} {unit}"
    if d >= 0.05:
        return f"rose by about {100 * d:.0f} {unit}"
    return "stayed roughly constant"


def build_texts(f, info):
    H = f["H"]
    t = {}
    per_pass = []
    for _, r in H.iterrows():
        per_pass.append(f"Passes {int(r.passes)}: best share {pc(r.best_frac)} CPU, {ms(r.best_ms)} versus GPU-only {ms(r.gpu_only_ms)} ({sx(r.speedup_best_vs_gpu_only)}) and CPU-only {ms(r.cpu_only_ms)} ({sx(r.speedup_best_vs_cpu_only)})")
    t["per_pass"] = per_pass

    h1_interior = f["interior"] >= 0.5 * f["n_cfg"]
    h1_trend = (not np.isnan(f["d_passes"])) and f["d_passes"] <= -0.05
    h2_rate = f["pred_in_range"] / f["n_cfg"]
    t["H1v"] = verdict(h1_interior and h1_trend, h1_interior or h1_trend)
    t["H2v"] = verdict(h2_rate >= 0.7, h2_rate >= 0.4)
    not_always = (f["beats_gpu"] < f["n_cfg"]) or (f["beats_single"] < f["n_cfg"])
    t["H3v"] = verdict(not_always)

    t["abstract"] = (
        f"A discrete GPU leaves the multicore CPU idle, but splitting a job across both devices adds copy and synchronisation costs. "
        f"This study measures how the share of work given to the CPU changes the makespan of an independent-task workload and where the optimum lies. "
        f"The workload is {nums(f['tiles_list'])} tiles of {f['tile_dim']}x{f['tile_dim']} floats, each filtered by {nums(f['passes_list'])} passes of a 3x3 blur. "
        f"A share of tiles from 0 to 100 percent runs on {nums(f['threads_list'])} OpenMP threads and the rest on {info.get('gpu_short', 'the GPU')} through CUDA. "
        f"Medians of {f['reps']} repetitions are compared with CPU-serial, CPU-only and GPU-only baselines. "
        f"At the largest size with {f['th_max']} threads the best split sent {join_and(pc(v) for v in H.best_frac)} of the tiles to the CPU "
        f"(passes {nums(H.passes)} respectively), a speedup of {join_and(sx(v) for v in H.speedup_best_vs_gpu_only)} over GPU-only. "
        f"Over all {f['n_cfg']} configurations the hybrid beat the faster single device by more than noise in {f['beats_single']}. "
        f"The optimal share {trend_word(f['d_passes'])} as passes rose from {f['p_min']} to {f['p_max']}, "
        f"and the rule f* = T_gpu / (T_cpu + T_gpu) landed in the near-optimal range in {f['pred_in_range']} of {f['n_cfg']} configurations."
    )
    return t


# ------------------------------------------------------------------ document
def build(args):
    res = Path(args.results)
    raw = pd.read_csv(res / "s03_raw_runs.csv")
    med = pd.read_csv(res / "s03_results_median.csv")
    opt = pd.read_csv(res / "s03_optimal_split.csv")
    sysinfo = read_sysinfo(res / "system_info.txt")
    figdir = res / "figures"
    hardware = str(raw.Hardware.iloc[0])
    if "NO GPU" in hardware and not args.allow_test_data:
        raise SystemExit("These results come from the host-only test build (no GPU). Run the real experiment on Kaggle with a GPU first.")
    f = compute_facts(raw, med, opt)
    gpu_line = sysinfo.get("GPU", "")
    gpu_name = gpu_line.split("|")[0].strip() if gpu_line else hardware.split("+")[0].strip()
    info = {"gpu_short": f"an {gpu_name}" if gpu_name[:1].upper() in "AEIOU" else f"a {gpu_name}"}
    t = build_texts(f, info)
    H, opt = f["H"], f["opt"]

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.left_margin = sec.right_margin = Inches(0.7)
    sec.top_margin, sec.bottom_margin = Inches(0.6), Inches(0.6)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(9.5)

    # title block
    p = para(doc, "**CS7 - S03: CPU-GPU Workload Partitioning**", size=14, after=0, align="c")
    para(doc, "BCSE412L Parallel Computing, Fall Semester 2026-27 | Individual Technical Report", size=9.5, after=1, align="c")
    para(doc, f"**Author:** {args.name}    **Registration No.:** {args.reg}    **Batch / Section:** {args.batch}    **Date:** {args.date}", size=9, after=3, align="c")

    heading(doc, "Abstract")
    para(doc, t["abstract"], align="j")

    heading(doc, "1. Introduction and Literature Survey")
    para(doc, (
        "A discrete GPU can run a data-parallel kernel far faster than a multicore CPU, but the data must cross the PCIe bus and the CPU sits idle while it waits. "
        "The CUDA Best Practices Guide therefore advises minimising host-device transfers and overlapping the remaining ones with computation [1]. "
        "Gregg and Hazelwood show that CPU-versus-GPU comparisons depend strongly on where the data resides and whether transfer time is counted [4], and Mittal and Vetter survey workload partitioning as one of the main techniques for using both processors [2]. "
        "Qilin automates the CPU/GPU mapping with an empirically trained performance model [3]. Nozal and Bosque extend oneAPI so that a single kernel co-executes on a CPU and an integrated GPU, and report that dynamic load balancing improves efficiency over static splits [5]. "
        "Li et al. split each neural-network layer by output channels between CPU and GPU and use latency predictors to choose the split that balances the two devices [6]. "
        "Recent work also models the cost, including energy, of CPU-GPU data transfer [7]."), align="j")
    para(doc, (
        "**Gap addressed.** These works study runtimes, predictors or one application. This study instead isolates one question on a discrete PCIe GPU: for independent tiles, how does the best static CPU share move with arithmetic intensity, "
        "CPU thread count and problem size, and is a rule built from two single-device timings enough to find it? It extends the CPU-versus-GPU crossover measured in S02 of this batch by examining the region between the two extremes."), align="j")

    heading(doc, "2. Problem, Research Question and Hypotheses")
    para(doc, (
        "**Problem.** A workload of N independent image tiles must be processed with the smallest makespan, using the CPU, the GPU, or both. "
        "**Research question (CS7 S03).** Is CPU+GPU execution always better than GPU-only execution, and what split minimises the makespan? "
        "**H1:** the optimal CPU share is neither 0 nor 100 percent and falls as arithmetic intensity (passes per tile) rises, because kernel work grows with passes while copy cost per tile does not. "
        "**H2:** the throughput-proportional share f* = T_gpu / (T_cpu + T_gpu), computed from the two single-device runs, predicts the measured optimum. "
        "**H3:** hybrid execution does not always beat GPU-only; where the CPU is slow relative to the GPU the gain disappears into noise."), align="j")

    heading(doc, "3. Methodology")
    cpu_line = sysinfo.get("CPU", "").split("|")[0].strip() or "CPU not recorded"
    nv = re.search(r"release ([\d.]+), (V[\d.]+)", sysinfo.get("CUDA compiler", ""))
    cuda_txt = f"nvcc {nv.group(1)} ({nv.group(2)})" if nv else "CUDA compiler version not recorded"
    clean = lambda s: s.replace(" | ", ", ")
    para(doc, (
        f"**Workload.** {nums(f['tiles_list'])} tiles of {f['tile_dim']}x{f['tile_dim']} single-precision pixels ({nums([v * f['tile_dim'] ** 2 * 4 / 2 ** 20 for v in f['tiles_list']])} MiB of input), "
        f"each processed by {nums(f['passes_list'])} passes of a 3x3 box blur with clamped edges inside the tile. Passes set the arithmetic intensity while the data volume stays fixed. Tiles are independent, so any subset can run on either device. "
        f"**Partition.** A share f of the tiles goes to the CPU (OpenMP, {nums(f['threads_list'])} threads, dynamic scheduling) and 1-f to the GPU (CUDA, {f['chunk']}-tile chunks: host-to-device copy, one kernel per pass, device-to-host copy, {f['pinned_txt']} host memory, {f['streams_phrase']}). "
        "The host enqueues the GPU work asynchronously, computes its own share, then waits for the GPU. f = 0 is the GPU-only baseline and f = 1 the CPU-only baseline; CPU-serial (1 thread, f = 1) is the reference for speedup. "
        f"CPU shares from 0 to 100 percent were swept in {f['step']}-point steps. "
        f"**Measurements.** Per run: makespan, enqueue (scheduling) time, CPU compute time and summed per-thread busy time, GPU copy and kernel times from CUDA events, and the host wait for the GPU (synchronisation). "
        f"Each point was repeated {f['reps']} times, alternating sweep direction; the median is reported and the spread is shown as min-max. GPU output was checked against the CPU result for every share (maximum relative error {f['err_max']:.1e}). "
        "Metrics: speedup = baseline median / makespan (baseline always named); partition efficiency = ideal makespan 1/(1/T_cpu + 1/T_gpu) divided by measured makespan; CPU utilisation = busy thread time / (threads x makespan); GPU kernel utilisation = kernel time / makespan. "
        f"**Environment.** GPU: {clean(sysinfo.get('GPU', gpu_name))}; CPU: {clean(sysinfo.get('CPU', cpu_line))}; RAM {sysinfo.get('RAM', 'n/a')}; {cuda_txt}; {sysinfo.get('Host compiler', 'n/a')}; OS {sysinfo.get('OS', 'n/a')}. "
        "Build: nvcc -O3 with host flags -fopenmp -O3 -march=native. Runtime: OMP_DYNAMIC=FALSE, CUDA_VISIBLE_DEVICES=0, no thread pinning. Source code, raw CSV and plotting scripts are in the submitted archive; its README lists the exact commands."), align="j")

    heading(doc, "4. Results and Discussion")
    rows = []
    for _, r in opt[(opt.threads_req == f["th_max"])].sort_values(["passes", "tiles"]).iterrows():
        rows.append([int(r.passes), int(r.tiles), f"{r.cpu_only_ms:.1f}", f"{r.gpu_only_ms:.1f}",
                     (pc(r.best_frac) if r.near_opt_lo == r.near_opt_hi else f"{pc(r.best_frac)} ({100 * r.near_opt_lo:.0f}-{100 * r.near_opt_hi:.0f}%)"), f"{r.best_ms:.1f}",
                     sx(r.speedup_best_vs_gpu_only), sx(r.speedup_best_vs_cpu_only), sx(r.speedup_best_vs_serial) if pd.notna(r.speedup_best_vs_serial) else "n/a",
                     pc(r.predicted_frac)])
    para(doc, f"**Table 1.** Summary at {f['th_max']} CPU threads (medians, times in ms). Best share = CPU share with the lowest median makespan; range = shares within 2% of it. Speedups name their baseline in the column header.", size=8, after=2, italic=False)
    table(doc, ["Passes", "Tiles", "CPU-only", "GPU-only", "Best CPU share (range)", "Best", "vs GPU-only", "vs CPU-only", "vs CPU serial", "Predicted share f*"], rows,
          widths=[0.5, 0.5, 0.65, 0.65, 1.2, 0.55, 0.7, 0.7, 0.7, 0.75])
    para(doc, "", size=3, after=2)

    doc.add_picture(str(figdir / "fig1_split_time_util_speedup.png"), width=Inches(6.7))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, "Figure 1. Makespan (with the CPU and GPU parts), utilisation and speedup against the share of tiles given to the CPU, at the largest size and thread count. Star = best share.")

    # headline paragraph
    s1 = "; ".join(t["per_pass"])
    para(doc, (
        f"**Headline result** ({f['t_max']} tiles, {f['th_max']} threads). {s1}. "
        f"Across all {f['n_cfg']} configurations the best hybrid split beat GPU-only by more than the noise threshold (twice the larger coefficient of variation, at least 2%) in {f['beats_gpu']} and beat the faster of the two single devices in {f['beats_single']}; "
        f"the speedup of the best split over GPU-only ranged from {sx(f['gain_min'])} to {sx(f['gain_max'])}. The GPU was the faster single device in {f['gpu_faster']} of {f['n_cfg']} configurations. "
        f"**Utilisation at the best split:** {f['util_txt']}."
        + (f" **Scalability and efficiency:** the speedup of the best split over CPU serial grew from {f['scal_txt']}; efficiency is speedup divided by (CPU threads + 1 GPU); it can exceed 1 because the GPU does far more work than one CPU thread, so it is not comparable with a pure CPU efficiency." if f["scal_txt"] else "")), align="j")

    doc.add_picture(str(figdir / "fig2_time_breakdown.png"), width=Inches(6.7))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, "Figure 2. Time decomposition at each CPU share. Orange: CPU compute wall time. Blue stack: GPU host-to-device copy, kernels and device-to-host copy. Black dash: makespan.")

    tf = H.set_index("passes").gpu_only_transfer_frac
    tf_lo, tf_hi = tf.loc[f["p_min"]], tf.loc[f["p_max"]]
    para(doc, (
        f"**Why the split behaves this way.** In the GPU-only runs, copies took {pc(tf_lo)} of the GPU time at passes {f['p_min']} and {pc(tf_hi)} at passes {f['p_max']}. Copy cost per tile is fixed while kernel cost grows with passes (Figure 2), "
        f"so the GPU becomes cheaper per tile as intensity rises. Measured over all configurations, the optimal CPU share {trend_word(f['d_passes'])} from passes {f['p_min']} to {f['p_max']}, "
        f"{trend_word(f['d_threads'])} from {f['th_min']} to {f['th_max']} threads, and {trend_word(f['d_size'])} from {min(f['tiles_list'])} to {f['t_max']} tiles. "
        + (f"At interior optima the CPU and GPU parts finished within {pc(f['imb_med'])} of the makespan of each other (median), " if not np.isnan(f["imb_med"]) else "")
        + f"and the best split reached a median {pc(f['pe_median'])} of the ideal perfect-overlap makespan. At the best split the host's final wait for the GPU was {pc(f['sync_rel'], 1)} and the enqueue (scheduling) time {pc(f['sched_rel'], 1)} of the makespan (medians)."), align="j")

    doc.add_picture(str(figdir / "fig3_optimal_split_trends.png"), width=Inches(6.7))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, "Figure 3. Optimal CPU share (solid, measured; dashed x, predicted f*) against thread count and problem size, and the gain of the best split over each single device.")

    rate_h2 = f["pred_in_range"] / f["n_cfg"]
    agree = "matched the grid optimum closely" if rate_h2 >= 0.7 else ("was only a rough guide" if rate_h2 >= 0.4 else "did not match the grid optimum")
    para(doc, (
        f"**Comparison with the literature.** The large copy share in the GPU-only runs ({pc(tf_lo)} of GPU time at passes {f['p_min']}) is the effect that the CUDA guide [1] and Gregg and Hazelwood [4] warn about. "
        f"Choosing the split so that both devices finish together is the idea behind the trained model of Qilin [3] and the latency-predictor balancing of [6]; here the simplest version of that idea, f*, {agree} ({f['pred_in_range']} of {f['n_cfg']} configurations). "
        "Nozal and Bosque [5] found dynamic balancing better than static splits; only static splits were tested here, so this study neither confirms nor contradicts that."), align="j")

    sb = f["single_best"]
    odd = ""
    if len(sb):
        ex = [f"passes {int(r.passes)}, {int(r.tiles)} tiles, {int(r.threads_req)} thread(s) ({'CPU-only' if r.best_frac == 1 else 'GPU-only'})" for _, r in sb.head(4).iterrows()]
        odd = f" **Unexpected or boundary results.** In {len(sb)} of {f['n_cfg']} configurations a single device was best, for example {join_and(ex)}; these are kept in the data and discussed as cases where the other device adds copy or scheduling cost without removing enough work."
    para(doc, (
        f"**Hypotheses.** H1 is {t['H1v']}: the optimum was interior in {f['interior']} of {f['n_cfg']} configurations and the optimal share {trend_word(f['d_passes'])} from passes {f['p_min']} to {f['p_max']}. "
        f"H2 is {t['H2v']}: f* fell inside the near-optimal range (within 2% of the best makespan, plus or minus 5 points) in {f['pred_in_range']} of {f['n_cfg']} configurations, with a mean absolute difference of {100 * f['pred_mae']:.0f} points from the grid optimum. "
        f"H3 is {t['H3v']}: the hybrid did not beat GPU-only by more than noise in {f['n_cfg'] - f['beats_gpu']} of {f['n_cfg']} configurations and did not beat the faster single device in {f['n_cfg'] - f['beats_single']}. "
        f"Run-to-run variation was a median {pc(f['cv_med'], 1)} (maximum {pc(f['cv_max'], 1)}) coefficient of variation of the makespan. The noisiest point (passes {f['noisy'][0]}, {f['noisy'][1]} tiles, {f['noisy'][2]} thread(s), {100 * f['noisy'][3]:.0f}% CPU share, median {ms(f['noisy'][4])}, CV {pc(f['noisy'][5], 1)}) was kept in the data and not removed.{odd}"), align="j")

    heading(doc, "5. Limitations and Threats to Validity")
    para(doc, (
        f"Only one GPU model, one CPU configuration and a shared cloud notebook were used, so the optimal shares are specific to this hardware (run-to-run variation up to {pc(f['cv_max'], 1)}). "
        "The workload is one regular, memory-light kernel; irregular kernels would shift the balance. The CPU kernel is a hand-vectorised OpenMP loop, not a tuned library, and the GPU kernel is a simple global-memory blur, so the CPU/GPU throughput ratio is implementation-specific. "
        f"Device buffers are allocated once and excluded from the timing; only the host-memory type and stream count listed in Section 3 were tested. Shares were swept on a {f['step']}-point grid, so optima are resolved to about {f['step']} points; the near-optimal range is reported for this reason. "
        "Thread counts above the physical core count share cores with the host thread that drives the GPU."), align="j")

    heading(doc, "6. Conclusion and Future Work")
    if f["beats_single"] == f["n_cfg"]:
        lead = f"For this tile-parallel workload CPU+GPU execution beat the faster single device by more than noise in every one of the {f['n_cfg']} configurations."
    elif f["beats_single"] == 0:
        lead = f"For this tile-parallel workload CPU+GPU execution never beat the faster single device by more than noise in the {f['n_cfg']} configurations."
    else:
        lead = f"For this tile-parallel workload CPU+GPU execution was not uniformly better: it beat the faster single device by more than noise in {f['beats_single']} of {f['n_cfg']} configurations."
    rate = f["pred_in_range"] / f["n_cfg"]
    rule = ("is a usable starting point" if rate >= 0.7 else ("is only a rough guide" if rate >= 0.4 else "did not locate the optimum reliably"))
    concl = (
        f"{lead} At {f['t_max']} tiles and {f['th_max']} threads the best split gave {join_and(sx(v) for v in H.speedup_best_vs_gpu_only)} over GPU-only (passes {nums(H.passes)}). "
        f"The throughput-proportional rule {rule} ({f['pred_in_range']} of {f['n_cfg']} configurations within the near-optimal range). "
        "Future work: (1) replace the static split with dynamic chunk scheduling or work stealing and compare it with the best static split; "
        "(2) overlap copies and kernels with several CUDA streams and measure how far the optimum moves; "
        "(3) calibrate f* online from the first few tiles, as adaptive-mapping systems do [3, 6]; "
        "(4) add energy to the objective, since transfers and idle devices change the energy-optimal split [7].")
    para(doc, concl, align="j")

    heading(doc, "References")
    for i, r in enumerate(REFS, 1):
        para(doc, f"[{i}] {r}", size=7.5, after=0.5)

    doc.save(args.out)
    print(f"saved {args.out}")
    wc = len(t["abstract"].split())
    print(f"abstract length: {wc} words (manual asks for 150-200)")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--results", default="results")
    ap.add_argument("--out", default="S03_Individual_Technical_Report.docx")
    ap.add_argument("--name", default="Mitisha Sharma")
    ap.add_argument("--reg", default="23BCE0434")
    ap.add_argument("--batch", default="________")
    ap.add_argument("--date", default="19 October 2026")
    ap.add_argument("--allow-test-data", action="store_true")
    build(ap.parse_args())
